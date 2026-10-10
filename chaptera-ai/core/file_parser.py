"""Extract text from TXT, DOCX, PDF.

DOCX bold runs are preserved as **bold** markers so preview/PDF/DOCX export can
reconstruct emphasis without requiring a rich-text editor.
"""

from pathlib import Path


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".txt":
        return data.decode("utf-8", errors="replace")
    if ext == ".doc":
        raise ValueError("Old .doc files are not supported. Save the Word file as .docx and upload again.")
    if ext == ".docx":
        return _docx_text(data)
    if ext == ".pdf":
        return _pdf_text(data)
    raise ValueError(f"Unsupported file type: {ext}. Use .txt, .docx, or .pdf")


def extract_metadata(filename: str, data: bytes) -> dict:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        return _docx_metadata(data)
    if ext == ".pdf":
        return _pdf_metadata(data)
    return {}


def _paragraph_has_page_break(paragraph) -> bool:
  from docx.oxml.ns import qn

  for br in paragraph._element.findall(".//" + qn("w:br")):
    if br.get(qn("w:type")) == "page":
      return True
  return False


def _docx_text(data: bytes) -> str:
    import io
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = _docx_paragraph_text(p).strip()
        if t:
            parts.append(t)
        if _paragraph_has_page_break(p):
            parts.append("\f")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n\n".join(parts)


def _docx_metadata(data: bytes) -> dict:
    import io
    from collections import Counter
    from docx import Document

    doc = Document(io.BytesIO(data))
    fonts: Counter[str] = Counter()
    sizes: Counter[int] = Counter()

    normal = doc.styles["Normal"].font
    if normal.name:
        fonts[normal.name] += 2
    if normal.size:
        sizes[round(normal.size.pt)] += 2

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.name:
                fonts[run.font.name] += len(run.text)
            if run.text and run.font.size:
                sizes[round(run.font.size.pt)] += len(run.text)

    meta: dict = {}
    if fonts:
        meta["body_font"] = fonts.most_common(1)[0][0]
    if sizes:
        meta["font_size"] = sizes.most_common(1)[0][0]
    return meta


def _docx_paragraph_text(paragraph) -> str:
    chunks: list[str] = []
    bold_open = False
    for run in paragraph.runs:
        text = run.text or ""
        if not text:
            continue
        is_bold = bool(run.bold)
        if is_bold and not bold_open:
            chunks.append("**")
            bold_open = True
        elif bold_open and not is_bold:
            chunks.append("**")
            bold_open = False
        chunks.append(text)
    if bold_open:
        chunks.append("**")
    return "".join(chunks) if chunks else paragraph.text


def _pdf_text(data: bytes) -> str:
    import io

    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError(
            "PDF support is not installed. Run: pip install pypdf"
        ) from e

    if not data:
        raise ValueError("PDF file is empty.")

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
    except Exception as e:
        raise ValueError(f"Could not read PDF: {e}") from e

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            pass
        if reader.is_encrypted:
            raise ValueError(
                "This PDF is password-protected. Remove the password and try again."
            )

    parts: list[str] = []
    for page in reader.pages:
        t = _extract_pdf_page_text(page)
        if t:
            parts.append(t)

    text = "\n\f\n".join(parts).strip()
    if not text:
        raise ValueError(
            "No readable text found in this PDF. "
            "It may be a scanned image (photo) PDF — export as DOCX or TXT, or use a PDF with selectable text."
        )
    return text


def _pdf_metadata(data: bytes) -> dict:
    import io

    try:
        from pypdf import PdfReader
    except ImportError:
        return {}

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        return {"page_count": len(reader.pages)}
    except Exception:
        return {}


def _extract_pdf_page_text(page) -> str:
    """Try default extraction, then layout mode for difficult PDFs."""
    for mode in (None, "layout"):
        try:
            if mode:
                t = page.extract_text(extraction_mode=mode)
            else:
                t = page.extract_text()
        except TypeError:
            t = page.extract_text()
        except Exception:
            t = None
        if t and t.strip():
            return t.strip()
    return ""
